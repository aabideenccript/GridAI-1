from typing import List, Dict, Any
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage
from langchain_openai import ChatOpenAI, OpenAIEmbeddings
from langchain_chroma import Chroma
from langchain.storage import InMemoryStore
from langchain.retrievers.multi_vector import MultiVectorRetriever
from langchain.retrievers import EnsembleRetriever
try:
    from langchain_community.retrievers import BM25Retriever
    BM25_AVAILABLE = True
except ImportError:
    BM25_AVAILABLE = False
import uuid
import logging
import pandas as pd
from docx import Document as DocxDocument
import json
import time

logger = logging.getLogger(__name__)

class RAGService:
    def __init__(self):
        self.vectorstore = None
        self.retriever = None
        self.ensemble_retriever = None
        self.store = InMemoryStore()
        self.model = ChatOpenAI(model="gpt-4o-mini", temperature=0.3, max_tokens=1500)
        self.conversations = {}
        self.all_documents = []
        self.stored_images = {}
        self.image_summaries = {}
    
    def process_docx(self, file_path: str) -> List[Document]:
        """Process DOCX files"""
        try:
            doc = DocxDocument(file_path)
            texts = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    texts.append(Document(page_content=para.text.strip()))
            
            for table_idx, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                
                if table_data:
                    table_json = json.dumps(table_data, indent=2)
                    texts.append(Document(page_content=table_json, metadata={"table": table_idx}))
            
            return texts
        except Exception as e:
            logger.error(f"Error processing DOCX: {str(e)}")
            return []
    
    def process_excel(self, file_path: str) -> List[Document]:
        """Process Excel files"""
        try:
            sheet_dataframes = pd.read_excel(file_path, sheet_name=None)
            texts = []
            
            for sheet_name, sheet_df in sheet_dataframes.items():
                table_json = sheet_df.to_json(orient='records', indent=2)
                texts.append(Document(page_content=table_json, metadata={"sheet": sheet_name}))
            
            return texts
        except Exception as e:
            logger.error(f"Error processing Excel: {str(e)}")
            return []
    
    def store_images(self, images: List[str], summaries: List[str]):
        """Store images with IDs for persistent access"""
        for i, (image, summary) in enumerate(zip(images, summaries)):
            img_id = f"img_{uuid.uuid4().hex[:8]}"
            self.stored_images[img_id] = {
                "data": image,
                "summary": summary,
                "timestamp": time.time()
            }
            self.image_summaries[img_id] = summary
        return list(self.stored_images.keys())
    
    def remove_image(self, img_id: str):
        """Remove image from storage and context"""
        if img_id in self.stored_images:
            del self.stored_images[img_id]
        if img_id in self.image_summaries:
            del self.image_summaries[img_id]
        logger.info(f"Removed image {img_id} from context")
    
    def get_relevant_images(self, query: str) -> List[Dict]:
        """Get images relevant to the query"""
        relevant_images = []
        query_lower = query.lower()
        
        for img_id, img_data in self.stored_images.items():
            summary_lower = img_data["summary"].lower()
            if any(word in summary_lower for word in query_lower.split() if len(word) > 3):
                relevant_images.append({
                    "id": img_id,
                    "data": img_data["data"],
                    "summary": img_data["summary"]
                })
        
        return relevant_images[:3]
    
    def setup_retriever(self, texts: List[Document], image_summaries: List[str], images: List[str] = None):
        """Setup advanced retriever with hybrid search"""
        try:
            self.vectorstore = Chroma(
                collection_name="electrical_rag",
                embedding_function=OpenAIEmbeddings(),
                persist_directory="./chroma_db"
            )
            
            self.retriever = MultiVectorRetriever(
                vectorstore=self.vectorstore,
                docstore=self.store,
                id_key="doc_id"
            )
            
            # Store all documents for hybrid search
            self.all_documents = []
            
            # Add text documents
            if texts:
                doc_ids = [str(uuid.uuid4()) for _ in texts]
                self.retriever.vectorstore.add_documents([
                    Document(page_content=doc.page_content, metadata={"doc_id": doc_ids[i]})
                    for i, doc in enumerate(texts)
                ])
                self.retriever.docstore.mset(list(zip(doc_ids, texts)))
                self.all_documents.extend(texts)
            
            # Add image summaries
            if image_summaries:
                if images:
                    self.store_images(images, image_summaries)
                
                img_ids = [str(uuid.uuid4()) for _ in image_summaries]
                summary_docs = [
                    Document(page_content=summary, metadata={"doc_id": img_ids[i]})
                    for i, summary in enumerate(image_summaries)
                ]
                self.retriever.vectorstore.add_documents(summary_docs)
                self.retriever.docstore.mset(list(zip(img_ids, image_summaries)))
                
                image_docs = [Document(page_content=summary) for summary in image_summaries]
                self.all_documents.extend(image_docs)
            
            # Setup hybrid retrieval
            self._setup_hybrid_retrieval()
            
            logger.info(f"RAG setup complete: {len(texts)} texts, {len(image_summaries)} images")
            
        except Exception as e:
            logger.error(f"Error setting up RAG: {str(e)}")
            raise
    
    def _setup_hybrid_retrieval(self):
        """Setup hybrid retrieval with BM25 and vector search"""
        try:
            if BM25_AVAILABLE and self.all_documents:
                # Setup BM25 retriever
                bm25_retriever = BM25Retriever.from_documents(self.all_documents)
                bm25_retriever.k = 5
                
                # Setup ensemble retriever
                vector_retriever = self.vectorstore.as_retriever(search_kwargs={"k": 5})
                
                self.ensemble_retriever = EnsembleRetriever(
                    retrievers=[bm25_retriever, vector_retriever],
                    weights=[0.4, 0.6]  # Favor vector search slightly
                )
                
                logger.info("Hybrid retrieval setup complete")
            else:
                logger.info("Using vector-only retrieval")
        except Exception as e:
            logger.error(f"Error setting up hybrid retrieval: {str(e)}")
    
    def _expand_query(self, query: str) -> List[str]:
        """Simple query expansion for better retrieval"""
        variations = [query]
        query_lower = query.lower()
        
        # Add electrical engineering related terms
        if any(term in query_lower for term in ['transformer', 'substation', 'power']):
            variations.append(query + " electrical equipment specifications")
        
        if any(term in query_lower for term in ['protection', 'relay', 'safety']):
            variations.append(query + " IEEE NESC standards")
        
        if any(term in query_lower for term in ['circuit', 'breaker', 'switch']):
            variations.append(query + " electrical protection systems")
        
        return variations[:3]  # Limit to 3 variations
    
    def _retrieve_with_expansion(self, query: str, k: int = 8) -> List[Document]:
        """Retrieve documents with query expansion"""
        try:
            expanded_queries = self._expand_query(query)
            all_docs = []
            doc_scores = {}
            
            for expanded_query in expanded_queries:
                if self.ensemble_retriever:
                    docs = self.ensemble_retriever.invoke(expanded_query)
                elif self.retriever:
                    docs = self.retriever.invoke(expanded_query)
                else:
                    docs = self.vectorstore.similarity_search(expanded_query, k=k//len(expanded_queries))
                
                # Score documents based on relevance
                for doc in docs:
                    doc_id = id(doc)
                    if doc_id not in doc_scores:
                        doc_scores[doc_id] = {"doc": doc, "score": 0}
                    doc_scores[doc_id]["score"] += 1
            
            # Sort by score and return top k
            sorted_docs = sorted(doc_scores.values(), key=lambda x: x["score"], reverse=True)
            return [item["doc"] for item in sorted_docs[:k]]
            
        except Exception as e:
            logger.error(f"Error in advanced retrieval: {str(e)}")
            # Fallback to simple retrieval
            if self.vectorstore:
                return self.vectorstore.similarity_search(query, k=k)
            return []
    
    async def query(self, message: str, session_id: str, context_images: List[str] = None) -> Dict[str, Any]:
        """Process query with advanced RAG and context images"""
        try:
            # Get conversation context
            context = self._get_context(session_id)
            
            # Check if user is referring to a previous image
            referring_to_image = any(phrase in message.lower() for phrase in 
                ['this image', 'that image', 'the image', 'based on that image', 'from the image', 
                 'tell me about this', 'tell me more', 'more about this', 'about this image',
                 'analyze this', 'what is this', 'describe this', 'about this iimage', 'this iimage'])
            
            logger.info(f"Referring to image: {referring_to_image}, Message: {message.lower()}")
            
            recent_image = None
            if referring_to_image or context_images:
                if not context_images:
                    recent_image = self._get_recent_image(session_id)
                    logger.info(f"Retrieved recent image: {recent_image is not None}")
                    if recent_image:
                        context_images = [recent_image]
                elif context_images:
                    logger.info(f"Using provided context images: {len(context_images)}")
                    referring_to_image = True  # Force image analysis if context images provided
            
            # Retrieve relevant documents with advanced methods
            docs = []
            if self.retriever:
                # Always try to retrieve relevant documents from uploaded content
                docs = self._retrieve_with_expansion(message, k=6)
                # If no relevant docs found, get some general docs
                if not docs:
                    docs = self.vectorstore.similarity_search("electrical", k=4)
                    if not docs:
                        docs = self.vectorstore.similarity_search("", k=4)
            
            # Build prompt with retrieved content
            doc_content = "\n\n".join([doc.page_content for doc in docs[:4]]) if docs else ""
            
            # Add context images if provided
            context_image_info = ""
            if context_images:
                if referring_to_image:
                    context_image_info = f"\n\nThe user is referring to a specific image they previously shared. Analyze this image directly to answer their question."
                else:
                    context_image_info = f"\n\nUser has selected {len(context_images)} images for context in this conversation."
            
            # Handle image-specific queries with vision analysis
            if context_images:
                logger.info(f"Analyzing image with vision model for query: {message}")
                # Use vision model to analyze the image directly
                image_analysis = await self._analyze_image_with_query(context_images[0], message)
                return {
                    "response": image_analysis,
                    "relevant_images": []
                }
            elif any(phrase in message.lower() for phrase in 
                  ['what is in', 'in this pdf', 'analyze this', 'what is here']) and doc_content:
                prompt = f"""
You are GridAI, an expert electrical engineering assistant. The user asked: "{message}"

Based on this uploaded electrical engineering content:
{doc_content}{context_image_info}

Provide a direct technical analysis describing the electrical components, systems, equipment ratings, and configuration shown. Be conversational and avoid repetitive sections.

Response:"""
            else:
                prompt = f"""
You are GridAI, an expert electrical engineering assistant with deep knowledge of power systems and grid infrastructure.

Conversation Context: {context}
Relevant Technical Documents: {doc_content}{context_image_info}
Current Query: {message}

Provide a direct, conversational response. Focus on answering the specific question with electrical engineering expertise.

Response:"""
            
            response = await self.model.ainvoke(prompt)
            
            # Store conversation with image context if present
            user_image = context_images[0] if context_images else None
            self._add_to_conversation(session_id, "user", message, user_image)
            self._add_to_conversation(session_id, "assistant", response.content)
            
            # Get relevant images for the query
            relevant_images = self.get_relevant_images(message)
            
            return {
                "response": response.content,
                "relevant_images": relevant_images
            }
            
        except Exception as e:
            logger.error(f"Error processing query: {str(e)}")
            return {
                "response": "I encountered an error processing your request. Please try again.",
                "relevant_images": []
            }
    
    def _get_context(self, session_id: str) -> str:
        """Get conversation context"""
        if session_id not in self.conversations:
            return ""
        
        recent = self.conversations[session_id][-4:]  # Last 4 messages
        return "\n".join([f"{msg['role']}: {msg['content']}" for msg in recent])
    
    def _get_recent_image(self, session_id: str) -> str:
        """Get the most recent image from conversation history"""
        if session_id not in self.conversations:
            return None
        
        # Look for the most recent image in conversation
        for msg in reversed(self.conversations[session_id]):
            if msg.get('image'):
                return msg['image']
        return None
    
    async def _analyze_image_with_query(self, image_base64: str, query: str) -> str:
        """Analyze image with specific query using vision model"""
        try:
            prompt = f"""
You are GridAI, an expert electrical engineering assistant. The user is asking: "{query}"

Analyze this electrical diagram/image and provide a direct, detailed response to their specific question. Focus on:
- What electrical components and systems you can see
- Technical specifications and configurations
- Power flow and system operation
- Any specific details relevant to their question

Be conversational and technical. Avoid generic boilerplate sections.
"""
            
            content = [
                {"type": "text", "text": prompt},
                {
                    "type": "image_url",
                    "image_url": {"url": f"data:image/jpeg;base64,{image_base64}"}
                }
            ]
            
            messages = [HumanMessage(content=content)]
            vision_model = ChatOpenAI(model="gpt-4o-mini", temperature=0.3, max_tokens=1000)
            response = await vision_model.ainvoke(messages)
            
            return response.content
            
        except Exception as e:
            logger.error(f"Error analyzing image with query: {str(e)}")
            return "I encountered an error analyzing the image. Please try again."
    
    def _add_to_conversation(self, session_id: str, role: str, content: str, image_data: str = None):
        """Add message to conversation with optional image context"""
        if session_id not in self.conversations:
            self.conversations[session_id] = []
        
        message = {
            "role": role,
            "content": content
        }
        
        if image_data:
            message["image"] = image_data
        
        self.conversations[session_id].append(message)
        
        # Keep only last 20 messages
        if len(self.conversations[session_id]) > 20:
            self.conversations[session_id] = self.conversations[session_id][-20:]